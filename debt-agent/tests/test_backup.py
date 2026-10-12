import io
import sqlite3
from datetime import datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import backup, db, export
from app import main as main_mod
from app.agent import Agent


def _count(path, table):
    c = sqlite3.connect(str(path))
    try:
        return c.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    finally:
        c.close()


def test_backup_is_a_full_working_copy(conn, tmp_path):
    cid = db.add_customer(conn, "أبو علي")
    db.add_transaction(conn, cid, "debt", 25000, "IQD")
    db.add_withdrawal(conn, 50000, "IQD", "مولدة")
    path = backup.make_backup(conn, tmp_path)
    assert path.name.endswith("-manual.db") and backup.NAME_RE.match(path.name)
    assert _count(path, "transactions") == 1 and _count(path, "withdrawals") == 1
    restored = db.connect(str(path))  # opens like the real notebook
    assert db.get_balance(restored, cid) == {"IQD": 25000, "USD": 0}
    restored.close()
    assert not list(tmp_path.glob("*.tmp"))


def test_daily_backup_once_a_day(conn, tmp_path):
    day = datetime(2026, 10, 12, 9, 0)
    first = backup.ensure_daily(conn, tmp_path, now=day)
    assert first.name == "debt-2026-10-12.db"
    assert backup.ensure_daily(conn, tmp_path, now=day.replace(hour=18)) is None  # same day: nothing new
    assert backup.ensure_daily(conn, tmp_path, now=day + timedelta(days=1)).name == "debt-2026-10-13.db"


def test_old_backups_are_removed(conn, tmp_path, monkeypatch):
    monkeypatch.setattr(backup, "KEEP_DAILY", 3)
    monkeypatch.setattr(backup, "KEEP_MANUAL", 2)
    start = datetime(2026, 10, 1, 9, 0)
    for i in range(5):
        backup.ensure_daily(conn, tmp_path, now=start + timedelta(days=i))
        backup.make_backup(conn, tmp_path, now=start + timedelta(days=i, hours=1))
    names = [b["name"] for b in backup.list_backups(tmp_path)]
    daily = [n for n in names if "manual" not in n]
    manual = [n for n in names if "manual" in n]
    assert daily == ["debt-2026-10-05.db", "debt-2026-10-04.db", "debt-2026-10-03.db"]  # newest kept
    assert len(manual) == 2 and manual[0].startswith("debt-2026-10-05")


def test_backup_path_only_backup_files(tmp_path):
    (tmp_path / "debt-2026-10-12.db").write_bytes(b"x")
    (tmp_path / ".env").write_text("SECRET")
    assert backup.backup_path("debt-2026-10-12.db", tmp_path)
    for bad in [".env", "../debt.db", "debt-2026-10-12.db/../../.env", "debt.db", ""]:
        assert backup.backup_path(bad, tmp_path) is None


def test_excel_export(conn):
    ali = db.add_customer(conn, "أبو علي")
    zain = db.add_customer(conn, "شركة الزين", kind="supplier")
    db.add_transaction(conn, ali, "debt", 25000, "IQD", source_message="=HYPERLINK(\"x\")")
    gone = db.add_transaction(conn, ali, "debt", 7000, "IQD")
    db.mark_undone(conn, gone)
    db.add_transaction(conn, zain, "debt", 1000000, "IQD")
    db.add_withdrawal(conn, 50000, "IQD", "مولدة")
    wb = load_workbook(io.BytesIO(export.build_workbook(conn)))
    assert wb.sheetnames == ["الزبائن (إلي)", "الموردين (عليّ)", "كل القيود", "السحوبات", "الجرار", "الوارد والصادر"]
    customers = list(wb["الزبائن (إلي)"].values)
    assert customers[1][:3] == ("أبو علي", 25000, 0) and customers[-1][:2] == ("المجموع", 25000)
    assert wb["الزبائن (إلي)"].sheet_view.rightToLeft
    entries = list(wb["كل القيود"].values)[1:]
    assert [e[3] for e in entries] == ["دين", "دين", "وارد"]
    assert entries[1][7] == "ملغي"  # undone entries listed and marked
    assert wb["كل القيود"]["G2"].data_type == "s"  # typed text, not a formula
    assert list(wb["السحوبات"].values)[1][1:4] == (50000, "دينار", "مولدة")


@pytest.fixture
def client(conn, tmp_path, monkeypatch):
    monkeypatch.setenv("BACKUP_DIR", str(tmp_path))
    monkeypatch.setattr(main_mod, "agent", Agent(conn, llm=object(), log_path=None))
    return TestClient(main_mod.app)


def test_backup_endpoints(client, conn, tmp_path):
    assert client.get("/backups").json()["backups"] == []
    r = client.post("/backups").json()
    assert r["ok"] and (tmp_path / r["name"]).is_file()
    listed = client.get("/backups").json()
    assert [b["name"] for b in listed["backups"]] == [r["name"]] and listed["folder"] == str(tmp_path)
    d = client.get("/backups/" + r["name"])
    assert d.status_code == 200 and d.content.startswith(b"SQLite format 3")
    assert client.get("/backups/.env").status_code == 404


def test_export_endpoint(client, conn):
    db.add_transaction(conn, db.add_customer(conn, "كرار"), "debt", 50, "USD")
    r = client.get("/export.xlsx")
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    assert list(load_workbook(io.BytesIO(r.content))["الزبائن (إلي)"].values)[1][:3] == ("كرار", 0, 50)
